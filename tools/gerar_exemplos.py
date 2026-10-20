"""Gera as 6 planilhas de exemplo dos painéis em tools/exemplos/, com dados 100% FICTÍCIOS.

Mesmas colunas e formatos das planilhas de verdade (docs/paineis.md), para testar o Hub e
apresentar os painéis sem expor números da empresa:

    .venv\\Scripts\\python -m tools.gerar_exemplos

O resultado é sempre o mesmo (sorteio com semente fixa): os testes contam com isso.
"""

import csv
import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from hub.paineis import FILES_BY_KEY

OUT = Path(__file__).parent / "exemplos"
MONTH_START = date(2026, 9, 1)
LAST_DAY = date(2026, 9, 17)  # uma quinta-feira: o mês "em andamento"
PASTAS = "ABCDEFG"
# quanto cada pasta já fez da meta do mês até o último dia (fictício, para os gráficos variarem)
PACE = {"A": .68, "B": .52, "C": .61, "D": .40, "E": .57, "F": .33, "G": .63, "TLV": .49}
# quanto da carteira (meta de positivação) cada pasta já positivou
POS_PACE = {"A": .62, "B": .48, "C": .71, "D": .39, "E": .55, "F": .45, "G": .66, "TLV": .52, "PLB": .58}
# meta financeira do mês por pasta (fictícia); TLV é um canal sem fornecedor próprio, como na planilha real
META_PASTA = {"A": 535000, "B": 618000, "C": 900000, "D": 615000, "E": 430000, "F": 417000, "G": 55000, "TLV": 30000}


def _xlsx(path: Path, header: list[str], rows: list[list]) -> None:
    book = Workbook()
    sheet = book.active
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="000000")
    for row in rows:
        sheet.append(row)
    book.save(path)


def generate(out: Path = OUT) -> None:
    rnd = random.Random(2026)
    out.mkdir(parents=True, exist_ok=True)
    name = lambda key, ext: out / (FILES_BY_KEY[key].base + ext)

    sellers = [(str(10 + i), f"VENDEDOR {i:02d}", f"SUPERVISOR {'XYZ'[i % 3]}") for i in range(1, 11)]
    # Televendas: aparece na 8081 como "P-TLV- ..." com o supervisor "P-TELEVENDAS (...)"
    sellers += [(str(60 + i), f"P-TLV- VENDEDOR {60 + i}", "P-TELEVENDAS (SUPERVISOR T)") for i in range(1, 3)]
    suppliers = []  # (id, nome, grupo, pasta)
    next_id = 1001
    for pasta in PASTAS:
        for n in range(1, 4):
            group = f"FORNECEDOR {pasta}{n}"
            parts = 2 if n == 1 else 1  # o primeiro de cada pasta tem duas divisões unificadas
            for k in range(parts):
                label = group + (f" DIV {k + 1}" if parts > 1 else "")
                suppliers.append((str(next_id), label, group + (" UNIFICADA" if parts > 1 else ""), pasta))
                next_id += 1
    _xlsx(name("fornecedores", ".xlsx"), ["ID FORNECEDOR", "FORNECEDOR", "FORNECEDOR II", "PASTA"],
          [[int(s[0]), s[1], s[2], s[3]] for s in suppliers])

    products = []  # (id, nome, fornecedor)
    for i in range(1, 161):
        supplier = suppliers[i % len(suppliers)]
        products.append((str(5000 + i), f"PRODUTO {i:03d}", supplier))
    focos = ["FOCO PILHAS", "FOCO SANDÁLIAS", "FOCO HIGIENE", "FOCO ESPONJAS", "FOCO CABELOS", "FOCO FARMA"]
    foco_of = {p[0]: focos[i % len(focos)] for i, p in enumerate(products[:60])}
    _xlsx(name("focos", ".xlsx"), ["ID PRODUTO", "PRODUTO", "PRODUTO FOCO"],
          [[int(p[0]), p[1], foco_of[p[0]]] for p in products[:60]])

    # meta financeira por fornecedor: a meta da pasta dividida entre os fornecedores dela
    by_pasta = {p: [s for s in suppliers if s[3] == p] for p in PASTAS}
    _xlsx(name("metas_forn", ".xlsx"), ["ID FORN", "NOME FORNECEDOR", "META FINANCEIRA"],
          [[int(s[0]), s[1], round(META_PASTA[s[3]] / len(by_pasta[s[3]]), -3)] for s in suppliers])

    kpis = []
    for sid, sname, sup in sellers:
        for i, foco in enumerate(focos):
            meta = rnd.randint(20, 36) + 12 * i
            # META = positivação (clientes); META II = financeiro (só em alguns focos, como na planilha real)
            kpis.append([int(sid), sname, sup, foco, "VAREJO" if i < 4 else "FARMA", meta,
                         round(rnd.uniform(3000, 7000) * (1 + i / 2), 2) if i % 2 == 0 else None])
    _xlsx(name("metas_kpi", ".xlsx"), ["ID", "VENDEDOR", "SUPERVISOR", "PRODUTO FOCO", "CLASSIFICACAO", "META",
                                       "META II"], kpis)

    # vendas: pedidos sorteados e depois escalados para cada pasta chegar no ritmo de PACE
    clients = [(str(200000 + i), f"CLIENTE {i:04d}") for i in range(1, 2501)]
    days = [MONTH_START + timedelta(n) for n in range((LAST_DAY - MONTH_START).days + 1)]
    lines = []
    order = 31000000
    for d in days:
        if d.weekday() == 6:
            continue
        for sid, sname, sup in sellers:
            for _ in range(rnd.randint(2, 5) if d.weekday() == 5 else rnd.randint(6, 12)):
                order += 1
                client = rnd.choice(clients)
                status = rnd.choice("FFFFFFLMB")  # cancelado não vem na extração
                for product in rnd.sample(products, rnd.randint(2, 6)):
                    qty = rnd.randint(1, 24)
                    lines.append([sid, sname, sup, order, status, d, client, product, qty, rnd.uniform(5, 40) * qty])
    group_of = lambda line: "TLV" if "TLV" in line[1] else line[7][2][3]  # a venda do Televendas é da TLV
    sold = {p: 0.0 for p in [*PASTAS, "TLV"]}
    reached = {p: set() for p in [*PASTAS, "TLV", "PLB"]}
    for line in lines:
        sold[group_of(line)] += line[9]
        reached[group_of(line)].add(line[6][0])
        reached["PLB"].add(line[6][0])
    factor = {p: META_PASTA[p] * PACE[p] / sold[p] for p in sold}
    carteira = {p: round(len(reached[p]) / POS_PACE[p]) for p in reached}
    _xlsx(name("totais", ".xlsx"), ["PASTA", "META", "CARTEIRA"],
          [[p, META_PASTA[p], carteira.get(p, 120)] for p in META_PASTA] + [["PLB", sum(META_PASTA.values()), carteira["PLB"]]])
    with open(name("vendas", ".csv"), "w", newline="", encoding="cp1252") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["ID RCA", "NOME RCA", "NOME SUPERVISOR", "NUMERO PEDIDO", "POSICAO PED", "DATA PEDIDO",
                    "ID CLIENTE", "CLIENTE", "ID PRODUTO", "NOME PRODUTO", "ID FORNECEDOR", "NOME FORNECEDOR",
                    "QTDE CX", "QTDE UND", "VALOR", "QTDE CX VENDIDA"])
        for sid, sname, sup, order, status, d, client, product, qty, value in lines:
            value *= factor["TLV" if "TLV" in sname else product[2][3]]
            br = lambda x: f"{x:.2f}".replace(".", ",")
            w.writerow([sid, sname, sup, order, status, d.strftime("%d/%m/%Y"), client[0], client[1], product[0],
                        product[1], product[2][0], product[2][1], 12, qty, br(value), br(qty / 12)])


if __name__ == "__main__":
    generate()
    print(f"Exemplos gravados em {OUT}")
