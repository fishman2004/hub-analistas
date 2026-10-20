"""Leitura das planilhas que os analistas exportam (CSV e XLSX), sem depender de como o Excel salvou.

- Cabeçalho: a primeira linha com conteúdo. Os nomes são comparados sem acento, sem diferença de
  maiúsculas e com espaços normalizados ("META POSITIVAÇÃO" = "Meta positivacao").
- CSV: aceita ";" (o padrão do Excel em português), "," ou tabulação, em UTF-8 ou no formato do
  Windows (cp1252).
- Números no formato brasileiro ("1.234,56") ou como o Excel guarda (1234.56).
- Datas "dd/mm/aaaa" (com ou sem hora) ou data de verdade do Excel.

Esta leitura serve a qualquer painel (fase 5) e ao histórico do Pombo (fase 9).
"""

import csv
import io
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path

EXTENSIONS = (".xlsx", ".xlsm", ".csv")


class PlanilhaError(Exception):
    """Arquivo que não dá para ler: a mensagem vai para a tela de status e para o hub.log."""


def normalize(name: object) -> str:
    text = unicodedata.normalize("NFKD", str(name or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", text).strip().upper()


def read_rows(path: Path, columns: dict[str, str]) -> list[dict]:
    """Lê só as colunas pedidas ({chave: nome da coluna}) e devolve uma lista de {chave: valor}.

    Falta de coluna é erro (com o nome do que faltou); colunas a mais são ignoradas, então nomes de
    clientes e vendedores nem chegam a ser guardados.
    """
    suffix = path.suffix.lower()
    if suffix == ".csv":
        rows = _csv_rows(path)
    elif suffix in (".xlsx", ".xlsm"):
        rows = _xlsx_rows(path)
    else:
        raise PlanilhaError(f"Formato {suffix or 'sem extensão'} não é aceito: use {', '.join(EXTENSIONS)}.")
    rows = iter(rows)
    header = None
    for row in rows:
        if any(v not in (None, "") for v in row):
            header = [normalize(v) for v in row]
            break
    if header is None:
        raise PlanilhaError("O arquivo está vazio.")
    positions = {}
    missing = []
    for key, column in columns.items():
        wanted = normalize(column)
        if wanted in header:
            positions[key] = header.index(wanted)
        else:
            missing.append(column)
    if missing:
        raise PlanilhaError("Faltam as colunas: " + ", ".join(missing) + ".")
    out = []
    for row in rows:  # continua de onde o cabeçalho parou
        if not any(v not in (None, "") for v in row):
            continue  # linha em branco no meio ou no fim
        out.append({key: (row[i] if i < len(row) else None) for key, i in positions.items()})
    return out


def _csv_rows(path: Path):
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    first = text.split("\n", 1)[0]
    delimiter = max(";,\t", key=first.count)
    return csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)


def _xlsx_rows(path: Path):
    from openpyxl import load_workbook  # só carrega quando há XLSX para ler

    try:
        book = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # zip quebrado, arquivo aberto e salvando pela metade etc.
        raise PlanilhaError(f"Não consegui abrir a planilha ({exc.__class__.__name__}).") from exc
    try:
        rows = [list(r) for r in book.worksheets[0].iter_rows(values_only=True)]
    finally:
        book.close()
    return rows


def number(value: object) -> float:
    """Número em qualquer formato que o Excel exporta. Vazio vale 0."""
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("R$", "").replace(" ", "")
    if not text:
        return 0.0
    if "," in text:
        text = text.replace(".", "").replace(",", ".")  # formato brasileiro: 1.234,56
    try:
        return float(text)
    except ValueError:
        raise PlanilhaError(f"Valor que não é número: {str(value)[:20]!r}.") from None


def code(value: object) -> str:
    """Código (de vendedor, cliente, produto...) como texto, sem o ".0" que o Excel às vezes põe."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text[:-2] if text.endswith(".0") and text[:-2].isdigit() else text


def day(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip().split(" ")[0]
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise PlanilhaError(f"Data que não entendi: {text[:20]!r} (use dd/mm/aaaa).")
