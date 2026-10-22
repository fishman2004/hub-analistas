"""PHD automático: a positivação de cada vendedor por dia, calculada da 8081 (fase 24).

Toda vez que a 8081 muda (pela pasta dos painéis ou pelo download do Drive), o Hub grava uma
planilha por pasta, "PHD - Pasta B.xlsx", no mesmo desenho do PHD que os analistas mandam no
grupo: título no topo, quadro por supervisor, linha 13 com o dia da semana, linha 14 com o dia do
mês, linha 15 com os nomes das colunas e o total de cada dia, e um vendedor por linha a partir da 16
(positivações de cada dia de F a AJ, média em AK). O Pombo de cada analista lê esse arquivo com o
modelo do PHD e manda um print com legenda por supervisor no horário.

Positivação do dia = clientes distintos que o vendedor atendeu naquele dia (todas as posições:
L, B, M e F; cancelado nem vem na extração).

A planilha tem nomes de vendedores: ela vai só para a pasta escolhida (a mesma pasta de rede do
histórico do Pombo, em geral), nunca para a TV.

A hora do arquivo gravado é a hora da 8081 que o gerou. Assim, se a 8081 parar de chegar, o
Pombo vê um PHD velho e não envia número de ontem como se fosse de hoje.
"""

import calendar
import json
import logging
import os
import tempfile
import threading
import time
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from .calendario import NO_HOLIDAYS, holidays
from .db import Database, get_setting, log_event, put_setting
from .paineis import FILES_BY_KEY, TLV, find_file, is_sales_day, panel_folder, seller_pasta
from .people import Person, ValidationError
from .planilhas import PlanilhaError, code, day, normalize, read_rows

log = logging.getLogger("hub.phd")

SETTING_FOLDER = "phd.pasta"  # vazio = "PHD" ao lado do banco do Hub
SETTING_AREAS = "phd.areas"  # uma linha por supervisor: "B-BJ = CAMPINA GRANDE"; "= FORA" tira do PHD
OUT = "FORA"
SETTING_STATE = "phd.estado"  # JSON: a 8081 usada, quando gerou, os arquivos e o último problema
DEFAULT_FOLDER_NAME = "PHD"
TICK_SECONDS = 60

COLUMNS = {"data": "DATA PEDIDO", "cliente": "ID CLIENTE", "rca": "NOME RCA", "supervisor": "NOME SUPERVISOR"}
ID_COLUMN = "ID RCA"  # opcional: sem ela, a coluna ID fica vazia

# o desenho da planilha (o mesmo do modelo do PHD no Pombo)
TITLE_ROW, SUMMARY_ROW, WEEKDAY_ROW, DAY_ROW, HEADER_ROW = 2, 7, 13, 14, 15
FIRST_DAY_COL = 6  # F
MEDIA_COL = FIRST_DAY_COL + 31  # AK
SUMMARY_MAX = WEEKDAY_ROW - SUMMARY_ROW - 2  # cabe o quadro de até 4 supervisores sem encostar na linha 13
WEEKDAYS = ["SEG", "TER", "QUA", "QUI", "SEX", "SÁB", "DOM"]
HEADERS = ("ÁREA", "ID", "VENDEDOR", "SUPERVISOR")

_lock = threading.Lock()  # uma geração por vez (o laço e o botão "Gerar agora")


@dataclass
class Vendedor:
    pasta: str
    supervisor: str
    nome: str
    rca_id: str = ""
    clientes: dict[int, set] = field(default_factory=dict)  # dia do mês -> clientes distintos

    def dia(self, n: int) -> int:
        return len(self.clientes.get(n, ()))


@dataclass(frozen=True)
class Mes:
    year: int
    month: int
    last: date  # último dia com venda no arquivo
    pastas: dict[str, list[Vendedor]]


# ---------- cálculo ----------

def read_8081(path: Path) -> list[dict]:
    try:
        return read_rows(path, {**COLUMNS, "id": ID_COLUMN})
    except PlanilhaError as exc:
        if normalize(ID_COLUMN) not in normalize(str(exc)):
            raise
        return read_rows(path, COLUMNS)


def compute(rows: list[dict]) -> Mes:
    """Positivação por vendedor e por dia no mês da venda mais recente, separada por pasta."""
    sales = []
    for r in rows:
        if r["data"] in (None, ""):
            continue  # linha de total sem data no fim da exportação
        sales.append((day(r["data"]), r))
    if not sales:
        raise PlanilhaError("Nenhuma venda no arquivo.")
    last = max(d for d, _ in sales)
    sellers: dict[tuple, Vendedor] = {}
    for d, r in sales:
        if (d.year, d.month) != (last.year, last.month):
            continue
        nome = str(r["rca"] or "").strip()
        supervisor = str(r["supervisor"] or "").strip()
        pasta = seller_pasta(nome, supervisor)
        if not nome or not pasta:
            continue
        key = (pasta, supervisor, nome)
        v = sellers.get(key)
        if v is None:
            v = sellers[key] = Vendedor(pasta, supervisor, nome, code(r.get("id")))
        v.clientes.setdefault(d.day, set()).add(code(r["cliente"]))
    pastas: dict[str, list[Vendedor]] = {}
    for v in sorted(sellers.values(), key=lambda v: (v.pasta, v.supervisor, v.nome)):
        pastas.setdefault(v.pasta, []).append(v)
    return Mes(last.year, last.month, last, pastas)


def reference_day(today: date) -> date:
    """O dia do resultado: ontem, e na segunda o sábado (domingo não tem venda)."""
    prev = today - timedelta(days=1)
    return prev - timedelta(days=1) if prev.weekday() == 6 else prev


def week_days(today: date, hol=NO_HOLIDAYS) -> list[date]:
    """Os dias da média: de segunda até o dia do resultado, sem domingo e sem feriado.

    É a conta da planilha dos analistas (=MÉDIASE(...;"<>DOM";...) com o intervalo da semana),
    sem precisar mudar o intervalo à mão a cada dia."""
    ref = reference_day(today)
    monday = ref - timedelta(days=ref.weekday())
    return [monday + timedelta(days=i) for i in range((ref - monday).days + 1)
            if is_sales_day(monday + timedelta(days=i), hol)]


def media(v: Vendedor, mes: Mes, today: date, hol=NO_HOLIDAYS) -> float:
    """Média de positivação por dia na semana do resultado (dias do mês da planilha)."""
    days = [d for d in week_days(today, hol) if (d.year, d.month) == (mes.year, mes.month)]
    return round(sum(v.dia(d.day) for d in days) / len(days), 2) if days else 0.0


def supervisor_code(supervisor: str) -> str:
    """"B-BJ (ALEXANDRE ...)" -> "B-BJ"."""
    return normalize(supervisor.split("(")[0]).strip(" -")


def parse_areas(text: str) -> dict[str, str]:
    out = {}
    for line in (text or "").splitlines():
        if "=" in line:
            key, value = line.split("=", 1)
            if key.strip() and value.strip():
                out[supervisor_code(key)] = value.strip().upper()
    return out


def file_name(pasta: str) -> str:
    return "PHD - Televendas.xlsx" if pasta == TLV else f"PHD - Pasta {pasta}.xlsx"


# ---------- a planilha ----------

def write_workbook(path: Path, pasta: str, mes: Mes, sellers: list[Vendedor], areas: dict[str, str],
                   today: date, hol=NO_HOLIDAYS) -> None:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    navy = PatternFill("solid", fgColor="0B2A3A")
    gray = PatternFill("solid", fgColor="595959")
    white = Font(bold=True, color="FFFFFF")
    center = Alignment(horizontal="center", vertical="center")
    days_in_month = calendar.monthrange(mes.year, mes.month)[1]
    day_cols = range(FIRST_DAY_COL, FIRST_DAY_COL + 31)

    wb = Workbook()
    ws = wb.active
    ws.title = "PHD"
    ws.cell(TITLE_ROW, 2, "PHD - POSITIVAÇÃO REPRESENTANTE DIA").font = Font(bold=True, size=20, color="FFFFFF")
    ws.cell(TITLE_ROW + 1, 2, f"Pasta {pasta} · {mes.month:02d}/{mes.year} · vendas até "
                              f"{mes.last.strftime('%d/%m')} · gerado pelo Hub a partir da 8081").font = Font(color="FFFFFF")
    for row in (TITLE_ROW, TITLE_ROW + 1):
        for c in range(2, MEDIA_COL + 1):
            ws.cell(row, c).fill = navy

    def area(v: Vendedor) -> str:
        return areas.get(supervisor_code(v.supervisor), supervisor_code(v.supervisor))

    def day_header(row: int) -> None:
        for i, c in enumerate(day_cols, start=1):
            if i <= days_in_month:
                ws.cell(row, c, WEEKDAYS[date(mes.year, mes.month, i).weekday()])
                ws.cell(row + 1, c, i)

    # quadro por supervisor (fica acima do cabeçalho; o Pombo lê só da linha 16 para baixo)
    supervisors = list(dict.fromkeys(v.supervisor for v in sellers))
    summary = len(supervisors) <= SUMMARY_MAX
    if summary:
        day_header(SUMMARY_ROW - 2)
        for i, h in enumerate(("ÁREA", "SUPERVISOR", "", "AREA"), start=2):
            ws.cell(SUMMARY_ROW, i, h or None)
        for n, sup in enumerate(supervisors, start=SUMMARY_ROW + 1):
            team = [v for v in sellers if v.supervisor == sup]
            ws.cell(n, 2, area(team[0]))
            ws.cell(n, 3, supervisor_code(sup))
            ws.cell(n, 5, sup)
            for i, c in enumerate(day_cols, start=1):
                if i <= days_in_month:
                    ws.cell(n, c, sum(v.dia(i) for v in team))
        for i, c in enumerate(day_cols, start=1):
            if i <= days_in_month:
                ws.cell(SUMMARY_ROW, c, sum(v.dia(i) for v in sellers))

    day_header(WEEKDAY_ROW)
    for i, h in enumerate(HEADERS, start=2):
        ws.cell(HEADER_ROW, i, h)
    for i, c in enumerate(day_cols, start=1):
        if i <= days_in_month:
            ws.cell(HEADER_ROW, c, sum(v.dia(i) for v in sellers))
    ws.cell(HEADER_ROW, MEDIA_COL, "MÉDIA")
    for n, v in enumerate(sellers, start=HEADER_ROW + 1):
        ws.cell(n, 2, area(v))
        ws.cell(n, 3, int(v.rca_id) if v.rca_id.isdigit() else (v.rca_id or None))
        ws.cell(n, 4, v.nome)
        ws.cell(n, 5, v.supervisor)
        for i, c in enumerate(day_cols, start=1):
            if i <= days_in_month:
                ws.cell(n, c, v.dia(i))
        ws.cell(n, MEDIA_COL, media(v, mes, today, hol)).number_format = "0.00"

    blocks = ((SUMMARY_ROW - 2, SUMMARY_ROW),) * summary + ((WEEKDAY_ROW, HEADER_ROW),)
    for rows in blocks:
        for r in range(rows[0], rows[1] + 1):
            for c in range(2, MEDIA_COL + 1):
                cell = ws.cell(r, c)
                if cell.value is not None or r == rows[1]:
                    cell.fill, cell.font, cell.alignment = (gray if r == rows[1] and c >= FIRST_DAY_COL else navy), white, center
    # as cores do PHD: 5 ou mais verde, 1 a 4 laranja, zero rosa (formatação condicional, como na planilha)
    first, last_col = get_column_letter(FIRST_DAY_COL), get_column_letter(FIRST_DAY_COL + 30)
    ranges = ((SUMMARY_ROW + 1, SUMMARY_ROW + len(supervisors)),) * summary + ((HEADER_ROW + 1, HEADER_ROW + len(sellers)),)
    for r0, r1 in ranges:
        if r1 < r0:
            continue
        rng = f"{first}{r0}:{last_col}{r1}"
        ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["5"],
                                                      fill=PatternFill("solid", fgColor="C6EFCE"), font=Font(color="1E6B34")))
        ws.conditional_formatting.add(rng, CellIsRule(operator="between", formula=["1", "4"],
                                                      fill=PatternFill("solid", fgColor="FCE4D6"), font=Font(color="C55A11")))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=["0"],
                                                      fill=PatternFill("solid", fgColor="FDE9E9"), font=Font(color="DE9696")))
    for col, width in (("A", 2), ("B", 18), ("C", 7), ("D", 40), ("E", 34)):
        ws.column_dimensions[col].width = width
    for c in list(day_cols) + [MEDIA_COL]:
        ws.column_dimensions[get_column_letter(c)].width = 5.5 if c != MEDIA_COL else 8
    ws.freeze_panes = ws.cell(HEADER_ROW + 1, FIRST_DAY_COL)

    # grava ao lado e troca de uma vez: o Pombo nunca lê a planilha pela metade
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(dir=path.parent, suffix=".tmp")
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    finally:
        Path(tmp).unlink(missing_ok=True)


# ---------- configuração e estado ----------

def default_folder(db: Database) -> Path:
    return db.path.parent / DEFAULT_FOLDER_NAME


def output_folder(db: Database) -> Path:
    custom = get_setting(db, SETTING_FOLDER)
    return Path(custom) if custom else default_folder(db)


def state(db: Database) -> dict:
    try:
        return json.loads(get_setting(db, SETTING_STATE) or "{}")
    except ValueError:
        return {}


def _save_state(db: Database, data: dict) -> None:
    with db.connect() as conn:
        put_setting(conn, SETTING_STATE, json.dumps(data, ensure_ascii=False))


def phd_config(db: Database) -> dict:
    s = state(db)
    folder = output_folder(db)

    def when(key):
        return datetime.fromisoformat(s[key]) if s.get(key) else None

    files = []
    for name in s.get("arquivos", []):
        p = folder / name
        files.append({"name": name, "exists": p.is_file()})
    return {"folder": get_setting(db, SETTING_FOLDER) or "", "default_folder": str(default_folder(db)),
            "current_folder": str(folder), "areas": get_setting(db, SETTING_AREAS) or "",
            "gerado": when("gerado"), "vendas_ate": when("vendas_ate"), "base": when("base"),
            "files": files, "error": s.get("erro", ""), "erro_em": when("erro_em")}


def save_phd_config(db: Database, now: datetime, actor: Person, folder: str, areas: str) -> None:
    if actor is None or not actor.is_admin:
        raise PermissionError("Só administradores podem fazer isso.")
    folder = folder.strip().strip('"')
    if folder and not Path(folder).is_dir():
        raise ValidationError(f"Não encontrei a pasta {folder}. Confira o caminho (ou deixe em branco para usar a padrão).")
    bad = [line.strip() for line in areas.splitlines() if line.strip() and "=" not in line]
    if bad:
        raise ValidationError(f"Use uma linha por supervisor no formato B-BJ = CAMPINA GRANDE (não entendi: {bad[0]}).")
    with db.connect() as conn:
        put_setting(conn, SETTING_FOLDER, folder or None)
        put_setting(conn, SETTING_AREAS, areas.strip() or None)
        log_event(conn, now, actor.id, "config_phd", "config:phd", "")


@dataclass(frozen=True)
class Resultado:
    ok: bool
    message: str
    files: tuple[str, ...] = ()


def _signature(db: Database, source: Path | None) -> str:
    if source is None:
        return ""
    st = source.stat()
    return json.dumps([str(source), st.st_mtime_ns, st.st_size, str(output_folder(db)),
                       get_setting(db, SETTING_AREAS) or ""])


def generate(db: Database, now: datetime, force: bool = False) -> Resultado:
    """Gera os PHDs se a 8081 (ou a configuração) mudou desde a última vez; com force, sempre."""
    with _lock:
        source = find_file(panel_folder(db), FILES_BY_KEY["vendas"].base)
        s = state(db)
        if source is None:
            return Resultado(False, "A 8081 ainda não está na pasta dos painéis.")
        signature = _signature(db, source)
        if not force and signature == s.get("assinatura") and not s.get("erro"):
            return Resultado(True, "O PHD já está em dia com a 8081.")
        try:
            mes = compute(read_8081(source))
            folder = output_folder(db)
            areas = parse_areas(get_setting(db, SETTING_AREAS) or "")
            hol = holidays(db)
            base_time = source.stat().st_mtime
            names = []
            for pasta, sellers in mes.pastas.items():
                sellers = [v for v in sellers if areas.get(supervisor_code(v.supervisor)) != OUT]
                if not sellers:
                    continue
                path = folder / file_name(pasta)
                write_workbook(path, pasta, mes, sellers, areas, now.date(), hol)
                os.utime(path, (base_time, base_time))  # a hora do PHD é a hora da 8081
                names.append(path.name)
        except (PlanilhaError, OSError) as exc:
            message = str(exc) if isinstance(exc, PlanilhaError) else (
                "Não consegui gravar o PHD na pasta (ela existe? algum PHD está aberto no Excel?).")
            log.warning("PHD não gerado: %s", exc)
            _save_state(db, {**s, "erro": message, "erro_em": now.isoformat(timespec="seconds")})
            return Resultado(False, message)
        _save_state(db, {"assinatura": signature, "gerado": now.isoformat(timespec="seconds"),
                         "vendas_ate": datetime.combine(mes.last, datetime.min.time()).isoformat(),
                         "base": datetime.fromtimestamp(base_time).isoformat(timespec="seconds"),
                         "arquivos": names, "erro": ""})
        return Resultado(True, f"PHD gerado para {len(names)} {'pasta' if len(names) == 1 else 'pastas'}.", tuple(names))


def loop(db_path: Path, tick_seconds: int = TICK_SECONDS) -> None:
    """Roda para sempre (linha de fundo do __main__): a cada minuto, se a 8081 mudou, gera o PHD."""
    while True:
        try:
            generate(Database(db_path), datetime.now())
        except Exception:
            log.exception("falha ao gerar o PHD")
        time.sleep(tick_seconds)
