"""Fase 24: o PHD (positivação representante dia) calculado da 8081, no desenho que o Pombo lê."""

import os
from datetime import date, datetime

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from conftest import at
from hub import phd
from hub.access import create_pin
from hub.calendario import NO_HOLIDAYS
from hub.paineis import panel_folder
from hub.people import ValidationError, add_person, get_person
from hub.web.app import create_app

BASE = "http://hub-sala:8800"
NOW = datetime(2026, 10, 22, 9, 0)  # quinta
HEADER = "ID RCA;NOME RCA;NOME SUPERVISOR;POSICAO PED;DATA PEDIDO;ID CLIENTE;VALOR"
# nomes inventados; o mesmo cliente em 2 linhas do mesmo dia conta 1 vez
LINES = [
    "11;B-LT-01- ANA FICTICIA;B-LT (MARINA COUTO);F;19/10/2026;100;10,00",
    "11;B-LT-01- ANA FICTICIA;B-LT (MARINA COUTO);L;19/10/2026;100;5,00",
    "11;B-LT-01- ANA FICTICIA;B-LT (MARINA COUTO);B;19/10/2026;101;5,00",
    "11;B-LT-01- ANA FICTICIA;B-LT (MARINA COUTO);M;20/10/2026;100;5,00",
    "11;B-LT-01- ANA FICTICIA;B-LT (MARINA COUTO);F;22/10/2026;102;5,00",
    "12;B-AG-01- BRUNO FICTICIO;B-AG (RAFAEL DUARTE);F;21/10/2026;200;5,00",
    "12;B-AG-01- BRUNO FICTICIO;B-AG (RAFAEL DUARTE);F;30/09/2026;201;5,00",  # outro mês: fora
    "13;C-XX-01- CAIO FICTICIO;C-XX (LUANA PRADO);F;21/10/2026;300;5,00",
    "14;P-TLV- DANI FICTICIA;P-TELEVENDAS (JOANA REIS);F;21/10/2026;400;5,00",
    "15;SEM PASTA;;F;21/10/2026;500;5,00",  # sem pasta: fora
    ";;;;;;55,00",  # linha de total no fim da exportação
]


def write_8081(folder, lines=LINES, header=HEADER, when=datetime(2026, 10, 22, 8, 30)):
    path = folder / "(DADOS) - 8081 - Vendas.csv"
    path.write_text("\n".join([header] + lines) + "\n", encoding="utf-8")
    os.utime(path, (when.timestamp(), when.timestamp()))
    return path


@pytest.fixture
def people(db):
    kayo = add_person(db, at(7), None, "Kayo")
    bia = add_person(db, at(7), None, "Bia")
    create_pin(db, at(7), kayo.id, "1234", "1234")  # administrador
    create_pin(db, at(7), bia.id, "1234", "1234")
    return get_person(db, kayo.id), get_person(db, bia.id), create_app(db, clock=lambda: NOW)


def login(app, person):
    c = TestClient(app, base_url=BASE)
    c.post(f"/entrar/{person.id}", data={"pin": "1234"})
    return c


# ---------- cálculo ----------

def test_positivacao_is_distinct_clients_per_seller_per_day(tmp_path):
    mes = phd.compute(phd.read_8081(write_8081(tmp_path)))
    assert (mes.year, mes.month, mes.last) == (2026, 10, date(2026, 10, 22))
    assert sorted(mes.pastas) == ["B", "C", "TLV"]
    bruno, ana = sorted(mes.pastas["B"], key=lambda v: v.nome)
    assert (ana.nome, ana.supervisor, ana.rca_id) == ("B-LT-01- ANA FICTICIA", "B-LT (MARINA COUTO)", "11")
    assert [ana.dia(d) for d in (19, 20, 21, 22)] == [2, 1, 0, 1]
    assert bruno.dia(21) == 1 and bruno.dia(30) == 0  # a venda de setembro ficou fora


def test_id_column_is_optional(tmp_path):
    header = HEADER.replace("ID RCA;", "")
    lines = [line.split(";", 1)[1] for line in LINES]
    mes = phd.compute(phd.read_8081(write_8081(tmp_path, lines, header)))
    assert mes.pastas["B"][0].rca_id == ""


def test_media_counts_sales_days_until_yesterday(tmp_path):
    mes = phd.compute(phd.read_8081(write_8081(tmp_path)))
    ana = next(v for v in mes.pastas["B"] if v.nome.startswith("B-LT"))
    # outubro/2026: de 1 a 21 há 18 dias de venda (seg a sáb); hoje (22) ainda não conta
    assert phd.media(ana, mes, date(2026, 10, 22)) == round(3 / 18, 2)

    class Feriado:
        def __contains__(self, d):
            return d == date(2026, 10, 12)
    assert phd.media(ana, mes, date(2026, 10, 22), Feriado()) == round(3 / 17, 2)


def test_areas_by_supervisor_code():
    areas = phd.parse_areas("b-lt = Litoral\n\nB-AG (RAFAEL DUARTE) = agreste\nlixo")
    assert areas == {"B-LT": "LITORAL", "B-AG": "AGRESTE"}
    assert phd.supervisor_code("B-LT (MARINA COUTO)") == "B-LT"
    assert phd.file_name("B") == "PHD - Pasta B.xlsx" and phd.file_name("TLV") == "PHD - Televendas.xlsx"


def test_workbook_has_the_phd_layout(tmp_path):
    mes = phd.compute(phd.read_8081(write_8081(tmp_path)))
    path = tmp_path / "PHD - Pasta B.xlsx"
    phd.write_workbook(path, "B", mes, mes.pastas["B"], {"B-LT": "LITORAL"}, date(2026, 10, 22), NO_HOLIDAYS)
    ws = load_workbook(path).active
    assert ws.title == "PHD" and ws["B2"].value.startswith("PHD - POSITIVAÇÃO")
    # o desenho do modelo do PHD no Pombo: dia da semana na 13, dia do mês na 14, cabeçalho na 15
    assert (ws["F13"].value, ws["F14"].value, ws["AJ14"].value) == ("QUI", 1, 31)
    assert [ws.cell(15, c).value for c in range(2, 6)] == ["ÁREA", "ID", "VENDEDOR", "SUPERVISOR"]
    assert ws["AK15"].value == "MÉDIA"
    col19 = 5 + 19  # F = dia 1
    assert ws.cell(15, col19).value == 2  # total do dia 19 na pasta
    rows = {ws.cell(r, 4).value: r for r in range(16, 18)}
    r = rows["B-LT-01- ANA FICTICIA"]
    assert (ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 5).value) == ("LITORAL", 11, "B-LT (MARINA COUTO)")
    assert [ws.cell(r, 5 + d).value for d in (19, 20, 21, 22)] == [2, 1, 0, 1]
    assert ws.cell(rows["B-AG-01- BRUNO FICTICIO"], 2).value == "B-AG"  # sem área cadastrada: o código
    assert ws.cell(18, 4).value is None
    # quadro por supervisor acima do cabeçalho
    assert {ws.cell(r, 3).value for r in (8, 9)} == {"B-LT", "B-AG"}
    assert ws.conditional_formatting  # as cores do PHD


def test_many_supervisors_skip_the_summary(tmp_path):
    lines = [f"{i};B-S{i}-01- V{i};B-S{i} (SUP {i});F;21/10/2026;{i};1,00" for i in range(1, 7)]
    mes = phd.compute(phd.read_8081(write_8081(tmp_path, lines)))
    path = tmp_path / "x.xlsx"
    phd.write_workbook(path, "B", mes, mes.pastas["B"], {}, date(2026, 10, 22))
    ws = load_workbook(path).active
    assert all(ws.cell(r, 3).value is None for r in range(7, 13))
    assert ws.cell(21, 4).value == "B-S6-01- V6"


# ---------- geração ----------

def test_generate_writes_one_file_per_pasta_with_the_8081_time(db, tmp_path):
    source = write_8081(panel_folder(db))
    result = phd.generate(db, NOW)
    assert result.ok and set(result.files) == {"PHD - Pasta B.xlsx", "PHD - Pasta C.xlsx", "PHD - Televendas.xlsx"}
    out = phd.output_folder(db) / "PHD - Pasta B.xlsx"
    assert out.stat().st_mtime == pytest.approx(source.stat().st_mtime)
    cfg = phd.phd_config(db)
    assert cfg["gerado"] == NOW and cfg["vendas_ate"].date() == date(2026, 10, 22) and not cfg["error"]
    assert phd.generate(db, NOW).message == "O PHD já está em dia com a 8081."
    write_8081(panel_folder(db), LINES[:3], when=datetime(2026, 10, 22, 9, 30))
    assert phd.generate(db, NOW).files == ("PHD - Pasta B.xlsx",)


def test_generate_without_8081_or_with_a_broken_one(db):
    assert not phd.generate(db, NOW).ok
    write_8081(panel_folder(db), [], header="OUTRA;COISA")
    result = phd.generate(db, NOW)
    assert not result.ok and "Faltam as colunas" in result.message
    assert "Faltam as colunas" in phd.phd_config(db)["error"]


def test_config_only_for_admin_and_checks_the_folder(db, people, tmp_path):
    kayo, bia, _ = people
    with pytest.raises(PermissionError):
        phd.save_phd_config(db, NOW, bia, "", "")
    with pytest.raises(ValidationError, match="Não encontrei a pasta"):
        phd.save_phd_config(db, NOW, kayo, str(tmp_path / "nao-existe"), "")
    with pytest.raises(ValidationError, match="B-BJ = CAMPINA GRANDE"):
        phd.save_phd_config(db, NOW, kayo, "", "B-LT LITORAL")
    phd.save_phd_config(db, NOW, kayo, f'"{tmp_path}"', "B-LT = LITORAL")
    assert phd.output_folder(db) == tmp_path


# ---------- página ----------

def test_panels_page_generates_and_downloads(db, people, tmp_path):
    kayo, bia, app = people
    write_8081(panel_folder(db))
    admin = login(app, kayo)
    assert "PHD automático" in admin.get("/paineis").text
    assert login(app, bia).post("/paineis/phd/gerar").status_code == 403
    assert "/paineis/phd/gerar" not in login(app, bia).get("/paineis").text
    r = admin.post("/paineis/phd/gerar")
    assert r.status_code == 200 and "PHD gerado para 3 pastas." in r.text and "PHD - Pasta B.xlsx" in r.text
    got = login(app, bia).get("/paineis/phd/PHD - Pasta B.xlsx")
    assert got.status_code == 200 and got.content[:2] == b"PK"
    assert admin.get("/paineis/phd/..%2Fhub.db").status_code == 404
    assert admin.get("/paineis/phd/PHD - Pasta Z.xlsx").status_code == 404
    assert TestClient(app, base_url=BASE).get("/paineis/phd/PHD - Pasta B.xlsx", follow_redirects=False).status_code in (302, 303)


def test_panels_page_config(db, people, tmp_path):
    kayo, _, app = people
    write_8081(panel_folder(db))
    admin = login(app, kayo)
    r = admin.post("/paineis/phd/configurar", data={"pasta": str(tmp_path), "areas": "B-LT = LITORAL"})
    assert "Configuração salva." in r.text and (tmp_path / "PHD - Pasta B.xlsx").is_file()
    r = admin.post("/paineis/phd/configurar", data={"pasta": "", "areas": "sem igual"})
    assert r.status_code == 422 and "B-BJ = CAMPINA GRANDE" in r.text
